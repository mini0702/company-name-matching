"""给研究者看的 xlsx 统一格式：第 1 行字段名（加粗），第 2 行字段说明（灰底小号斜体、自动换行，注明来源，如"官网域名【样本·Crunchbase】"），
研究者要填的列标黄并放在前面；长文本列加宽换行；冻结在第 2 行下面。

用法（在程序里）：
  from xlsx_style import write_xlsx, restyle
  write_xlsx("结果.xlsx", {"说明": ["第一句说明", "第二句"], "人工判断": (df, DESC)}, fill_cols=("你的判断",))
  restyle("已有文件.xlsx")      # 只改格式，不改内容（研究者填过的内容保留）
DESC 是 {列名: 中文说明}；没写说明的列第 2 行留空（会提示）。
"""
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

LONG = ("description", "reason", "evidence", "简介", "理由", "依据", "说明", "question", "问题")
MID = ("name", "company", "url", "domain", "linkedin", "parent", "acquirer", "decision", "名称", "网址", "域名", "母公司", "收购方", "结论")
FILL_DEFAULT = ("你的判断", "你的裁定", "审核", "审核理由")
GRAY, YELLOW = "F2F2F2", "FFF2CC"


def _clean(df):
    return df.apply(lambda s: s.map(lambda x: ILLEGAL_CHARACTERS_RE.sub("", x) if isinstance(x, str) else x))


def with_desc(df, desc):
    """在数据上方加一行字段说明。"""
    import pandas as pd
    miss = [c for c in df.columns if c not in desc]
    if miss:
        print("提示：这些列没有写第 2 行说明：", miss)
    return pd.concat([pd.DataFrame([{c: desc.get(c, "") for c in df.columns}]), _clean(df).astype(object)], ignore_index=True)


def style(ws, fill_cols=FILL_DEFAULT, desc_row=True, freeze=None):
    head = [c.value for c in ws[1]]
    first_data = 3 if desc_row else 2
    for j, name in enumerate(head, start=1):
        col = ws.cell(row=1, column=j).column_letter
        key = str(name or "").lower()
        long = any(k.lower() in key for k in LONG)
        width = 45 if long else 24 if any(k.lower() in key for k in MID) else 13
        if name in fill_cols:
            width = 30 if "理由" in str(name) else 14
        ws.column_dimensions[col].width = width
        h = ws.cell(row=1, column=j)
        h.font = Font(bold=True)
        if name in fill_cols:
            h.fill = PatternFill("solid", fgColor=YELLOW)
        if desc_row and ws.max_row >= 2:
            c = ws.cell(row=2, column=j)
            c.font = Font(italic=True, color="555555", size=9)
            c.fill = PatternFill("solid", fgColor=YELLOW if name in fill_cols else GRAY)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if long:
            for r in range(first_data, ws.max_row + 1):
                ws.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    if desc_row:
        ws.row_dimensions[2].height = 75
    n_fix = 1 + sum(1 for x in head[:3] if x in fill_cols)
    ws.freeze_panes = freeze or ws.cell(row=first_data, column=n_fix + 1).coordinate


def style_note(ws, width=120):
    ws.column_dimensions["A"].width = width
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(path, sheets, fill_cols=FILL_DEFAULT):
    """sheets：{表名: 说明句子列表 | DataFrame | (DataFrame, DESC)}。有 DESC 的表加第 2 行说明。"""
    import pandas as pd
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, obj in sheets.items():
            if isinstance(obj, list):
                pd.DataFrame({"说明": obj}).to_excel(w, sheet_name=name, index=False)
                style_note(w.sheets[name])
            elif isinstance(obj, tuple):
                with_desc(*obj).to_excel(w, sheet_name=name, index=False)
                style(w.sheets[name], fill_cols)
            else:
                _clean(obj).to_excel(w, sheet_name=name, index=False)
                style(w.sheets[name], fill_cols, desc_row=False)


def restyle(path, fill_cols=FILL_DEFAULT, note_sheets=("说明",)):
    """给已有 xlsx 套格式：第 2 行像说明（含"【"或"编号"）的按有说明行处理。"""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        v = ws.cell(row=2, column=1).value
        if ws.title in note_sheets or ws.max_column == 1:
            style_note(ws)
        elif isinstance(v, str) and ("【" in v or "编号" in v):
            style(ws, fill_cols)
        else:
            style(ws, fill_cols, desc_row=False)
    wb.save(path)


if __name__ == "__main__":
    import sys
    for p in sys.argv[1:]:
        restyle(p)
        print("已套格式", p)
